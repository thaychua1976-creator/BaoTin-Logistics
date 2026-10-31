import streamlit as st
import pandas as pd
import datetime, time, re
import tempfile, os, io
import urllib.request

from utils_core import parse_money_input, tao_tieu_de_kem_nut_refresh
from co_manager import save_co_transaction, delete_co_transaction, get_don_gia_co_theo_khach_hang
from dotenv import load_dotenv
# Vẫn load file .env cho môi trường Local
load_dotenv(override=False) 

# --- HỆ THỐNG CACHE BỘ NHỚ ĐỆM ---
@st.cache_data(ttl=1800, show_spinner=False)
def get_cached_master_data(_db_instance, query, params=None):
    return _db_instance.execute_query(query, params)

def clear_master_cache():
    get_cached_master_data.clear()
# ---------------------------------

# 1. KIỂM TRA ĐĂNG NHẬP (Bắt buộc)
if 'username' not in st.session_state or not st.session_state['username']:
    st.error("⚠️ Phiên đăng nhập đã hết hạn hoặc bạn chưa đăng nhập. Vui lòng đăng nhập lại!")
    st.stop() # Dừng toàn bộ code phía dưới nếu chưa đăng nhập

# 2. LẤY CHÍNH XÁC USERNAME ĐÃ LOGIN
current_user = st.session_state['username']

# 3. KIỂM TRA KẾT NỐI DATABASE
db = st.session_state.get('db')
if not db:
    st.error("⚠️ Lỗi kết nối Cơ sở dữ liệu.")
    st.stop()

def get_idx(lst, val, default=0):
    try: return lst.index(val)
    except: return default

st.markdown("<h3 style='text-align: center; color: #0b5394;'>📄 PHÂN HỆ QUẢN LÝ CHỨNG TỪ C/O (XUẤT KHẨU)</h3>", unsafe_allow_html=True)
st.toast("✅ Đang load số liệu bạn vui lòng đợi nha:")
st.divider()

# Thêm tab OCR vào danh sách tabs
tab_ocr_co, = st.tab(["🤖 TRÍCH XUẤT C/O TẠO FILE GIAO ĐẠI LÝ"])
  


# ==========================================
# TAB 3: TRÍCH XUẤT C/O BẰNG AI & ĐẠI LÝ KÝ NHẬN BẢN GỐC
# ==========================================
with tab_ocr_co:
    st.markdown("#### 🤖 Trích Xuất Dữ Liệu C/O & Ký Nhận Bàn Giao")
    
    # 🚀 1. SỬ DỤNG @st.fragment ĐỂ CÔ LẬP RERUN (Tuyệt đối không giật trang khi gõ/ký tên)
    @st.fragment
    def vung_trich_xuat_va_ky_nhan():
        # --- BƯỚC 1: TIÊM CSS ÉP KHOẢNG CÁCH NGAY TỪ ĐẦU ---
        st.markdown("""
            <style>
                div[data-testid="stVerticalBlock"] > div { padding-bottom: 0.1rem !important; }
                div[data-testid="stHorizontalBlock"] > div { padding-bottom: 0.1rem !important; }
                .stTextInput > div, .stDateInput > div { margin-bottom: -15px; }
                h4, h5 { margin-top: 5px !important; margin-bottom: 5px !important; }
                div[role="radiogroup"] { margin-bottom: -10px !important; margin-top: -10px !important; }
                div[data-testid="stWebComponentContainer"] { touch-action: none !important; overscroll-behavior: none !important; }
                iframe[title="streamlit_drawable_canvas.st_canvas"] { touch-action: none !important; }
            </style>
        """, unsafe_allow_html=True)

        try:
            from streamlit_drawable_canvas import st_canvas
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.drawing.image import Image as OpenpyxlImage
            from PIL import Image
            import pdfplumber
            import easyocr
            import numpy as np
        except ImportError:
            st.error("⚠️ Server thiếu thư viện. Chạy lệnh: `pip install streamlit-drawable-canvas openpyxl pdfplumber easyocr pypdf fpdf2`")
            st.stop()

        # [CẬP NHẬT 1]: THÊM OPTION 3 LÊN TRÊN CÙNG
        che_do_tab3 = st.radio(
            "Lựa chọn luồng công việc:", 
            ["1️⃣ Quét PDF Mới & Chỉnh Sửa", "2️⃣ Tải File Excel Cũ Lên (Làm tiếp)", "3️⃣ Ghép Ảnh C/O Vào File Đã Ký"],
            horizontal=True
        )
        
        st.markdown("<hr style='margin: 10px 0px 5px 0px; border: 1px solid #e0e0e0;'>", unsafe_allow_html=True)
        
        if "ocr_data" not in st.session_state:
            st.session_state["ocr_data"] = []
        if "reset_uploader_key" not in st.session_state:
            st.session_state["reset_uploader_key"] = 0
        
        # ---------------------------------------------------------
        # CHẾ ĐỘ 1: QUÉT PDF MỚI VÀ CHỈNH SỬA TRỰC TIẾP
        # ---------------------------------------------------------
        if che_do_tab3 == "1️⃣ Quét PDF Mới & Chỉnh Sửa":
            current_pdf_key = f"upload_pdfs_ocr_{st.session_state['reset_uploader_key']}"
            uploaded_pdfs = st.file_uploader("📂 Kéo thả file PDF C/O vào đây", type=["pdf"], accept_multiple_files=True, key=current_pdf_key)
            
            if uploaded_pdfs:
                if st.button("🚀 Bắt Đầu Quét & Trích Xuất", type="primary"):
                    with st.spinner("⏳ Đang xử lý siêu tốc..."):
                        extracted_data = []
                        pattern = r'([A-Z0-9.\-\s]{0,20})(\d{2})\s*[-/|17lI\.]\s*(\d{2})\s*[-/|17lI\.]\s*([0-9A-Z\s]{5,15})'
                        
                        for pdf_file in uploaded_pdfs:
                            try:
                                with pdfplumber.open(pdf_file) as pdf:
                                    for page in pdf.pages:
                                        co_number = "Không nhận diện được"
                                        text_fast = (page.extract_text() or "").upper()
                                        match = re.search(pattern, text_fast)
                                        
                                        if not match:
                                            bounding_box = (page.width * 0.35, 0, page.width, page.height * 0.35)
                                            micro_crop = page.crop(bounding_box)
                                            pil_img = micro_crop.to_image(resolution=300).original
                                            
                                            # FIX 1: Khởi tạo api_key từ trước để tránh UnboundLocalError
                                            api_key = None 
                                            try:
                                                # FIX 2: Sửa lại đúng tên thư viện theo cú pháp bên dưới
                                                import google.generativeai as genai
                                                if "api_keys" in st.secrets and "gemini" in st.secrets["api_keys"]:
                                                    api_key = st.secrets["api_keys"]["gemini"]
                                                elif "GEMINI_API_KEY" in st.secrets:
                                                    api_key = st.secrets["GEMINI_API_KEY"]
                                            except Exception:
                                                pass
                                                
                                            if not api_key:
                                                api_key = os.getenv("GEMINI_API_KEY")
                                                
                                            if api_key:
                                                genai.configure(api_key=api_key)
                                                model = genai.GenerativeModel('gemini-3.6-flash')
                                                prompt = "Hãy tìm và trích xuất duy nhất dòng chứa mã số C/O (Certificate of Origin) trong hình ảnh. Mã này thường có chứa số tham chiếu, mã quốc gia hoặc ngày tháng (ví dụ: VN-... hoặc dãy số đặc trưng ở góc trên). Chỉ trả về chuỗi mã số đó, tuyệt đối không trích xuất các văn bản khác và không giải thích thêm."
                                                response = model.generate_content([prompt, pil_img])
                                                text_ai = (response.text or "").upper()
                                                match = re.search(pattern, text_ai)
                                            else:
                                                # FIX 3: Dùng break thay vì return None để không làm sập toàn bộ fragment
                                                st.error("Lỗi: Không tìm thấy GEMINI_API_KEY trong hệ thống!")
                                                break    
                                            
                                        if match:
                                            prefix_raw, dd, mm = match.group(1), match.group(2), match.group(3)
                                            raw_suffix = match.group(4)
                                            num_clean = re.sub(r'\s+', '', raw_suffix).replace('O', '0')
                                            
                                            m_strict = re.match(r'^(\d{5,8})([A-Z]?)', num_clean)
                                            if m_strict:
                                                digits = m_strict.group(1)
                                                letter = m_strict.group(2)
                                                if len(digits) >= 7 and digits[6] == '8':
                                                    digits = digits[:6] 
                                                    letter = 'B'        
                                                elif len(digits) > 6:
                                                    digits = digits[:6]
                                                if letter != 'B':
                                                    letter = ''
                                                num_clean = digits + letter
                                                
                                            suffix_clean = f"{dd}/{mm}/{num_clean}"
                                            prefix_clean = re.sub(r'[\s.\-]', '', prefix_raw)
                                            for noise in ['EUR1', 'EUR', 'NO', 'NUM']:
                                                prefix_clean = prefix_clean.replace(noise, '')
                                                
                                            country = "??"
                                            for c in ['DE', 'NL', 'SE', 'UK', 'KZ', 'K2', 'K7']:
                                                if c in prefix_clean:
                                                    country = c.replace('K2', 'KZ').replace('K7', 'KZ')
                                                    break
                                            if country == "??" and len(prefix_clean) >= 2: country = prefix_clean[-2:]
                                            if country == "??" and prefix_clean.endswith('E'): country = 'DE'
                                            co_number = f"VN-{country} {suffix_clean}" if country != "??" else f"VN-?? {suffix_clean}"
                                        
                                        extracted_data.append({"STT": len(extracted_data) + 1, "SỐ C/O BẢN GỐC": co_number, "KÝ NHẬN": "", "NGÀY NHẬN": ""})
                            except Exception as inner_e:
                                st.warning(f"Lỗi đọc {pdf_file.name}: {inner_e}")
                        
                        # FIX 4: Chỉ gọi rerun() nếu có data. Nếu mảng rỗng thì giữ lại màn hình để user kịp đọc thông báo lỗi.
                        if extracted_data:
                            st.session_state["ocr_data"] = extracted_data
                            st.rerun()

        # ---------------------------------------------------------
        # CHẾ ĐỘ 2: TẢI LÊN FILE EXCEL (KHÔI PHỤC PHIÊN LÀM VIỆC)
        # ---------------------------------------------------------
        elif che_do_tab3 == "2️⃣ Tải File Excel Cũ Lên (Làm tiếp)":
            current_excel_key = f"upload_excel_to_sign_{st.session_state['reset_uploader_key']}"
            uploaded_excel = st.file_uploader("📂 Kéo thả file Excel (Bản chưa ký) vào đây", type=["xlsx", "xls"], key=current_excel_key)
            if uploaded_excel:
                try:
                    df_load = pd.read_excel(uploaded_excel, skiprows=3)
                    if 'SỐ C/O BẢN GỐC' in df_load.columns:
                        st.session_state["ocr_data"] = df_load.to_dict('records')
                        st.session_state["ten_file_goc"] = uploaded_excel.name
                        st.success("✅ Đã khôi phục dữ liệu từ file Excel thành công!")
                    else:
                        st.error("❌ File Excel không đúng định dạng chuẩn của phần mềm.")
                except Exception as e:
                    st.error(f"Lỗi đọc file Excel: {e}")

        # ---------------------------------------------------------
        # CHẾ ĐỘ 3: GHÉP ẢNH BỔ SUNG VÀO FILE PDF ĐÃ KÝ
        # ---------------------------------------------------------
        elif che_do_tab3 == "3️⃣ Ghép Ảnh C/O Vào File Đã Ký":
            st.info("💡 Mẹo: Ở công ty/cảng anh chỉ cần cho đại lý ký rồi tải file PDF chứa chữ ký về máy. Sau đó về nhà có mạng Wifi mạnh, anh tải ngược file PDF đó lên đây cùng với mớ hình ảnh C/O để hệ thống tự ghép ảnh vào đuôi PDF nhé!")
            
            col_up1, col_up2 = st.columns(2)
            
            # [CẬP NHẬT 1] - Đổi key tĩnh thành key động để dễ dàng xóa trắng
            current_pdf_key = f"upload_pdf_da_ky_{st.session_state['reset_uploader_key']}"
            current_img_key = f"upload_img_bo_sung_{st.session_state['reset_uploader_key']}"
            
            file_pdf_da_ky = col_up1.file_uploader("1️⃣ Tải lên Phiếu bàn giao ĐÃ KÝ (PDF)", type=["pdf"], key=current_pdf_key)
            imgs_bo_sung = col_up2.file_uploader("2️⃣ Tải lên hình ảnh C/O gốc bổ sung", type=["jpg", "jpeg", "png"], accept_multiple_files=True, key=current_img_key)
            
            if st.button("⚙️ BẮT ĐẦU GHÉP ẢNH VÀO PDF", type="primary"):
                if not file_pdf_da_ky:
                    st.error("⚠️ Vui lòng tải lên file PDF Phiếu bàn giao đã ký.")
                elif not imgs_bo_sung:
                    st.error("⚠️ Vui lòng tải lên ít nhất 1 hình ảnh C/O.")
                else:
                    with st.spinner("⏳ Đang tính toán không gian và nối ảnh, vui lòng chờ trong giây lát..."):
                        try:
                            import pypdf
                            from fpdf import FPDF
                            from PIL import Image as PILImage
                            import tempfile, os, io
                            import pdfplumber
                            
                            # --- 1. PHÂN TÍCH KHÔNG GIAN TRỐNG CỦA TRANG CUỐI BẰNG PDFPLUMBER ---
                            file_bytes = file_pdf_da_ky.getvalue()
                            lowest_y_mm = 0
                            try:
                                with pdfplumber.open(io.BytesIO(file_bytes)) as p_pdf:
                                    last_page = p_pdf.pages[-1]
                                    # Lấy tất cả đối tượng (chữ, hình ảnh chữ ký, đường kẻ) trên trang cuối
                                    all_objs = last_page.chars + last_page.images + last_page.rects + last_page.curves
                                    if all_objs:
                                        # Tìm tọa độ điểm thấp nhất (tính từ trên xuống dưới theo point)
                                        max_bottom_pt = max([obj['bottom'] for obj in all_objs if 'bottom' in obj], default=0)
                                        lowest_y_mm = max_bottom_pt / 2.83465 # Đổi từ point (72dpi) sang mm
                            except Exception as e:
                                print("Lỗi quét không gian PDF:", e)

                            # Chiều cao A4 ngang là 210mm. Một hàng ảnh cao tối đa 75mm + Tiêu đề = ~85mm
                            # Nếu điểm kết thúc nhỏ hơn 125mm -> Đủ không gian để chèn nối tiếp cùng trang
                            if 0 < lowest_y_mm < 125: 
                                start_y = lowest_y_mm + 15
                                is_overlay = True
                            else:
                                start_y = 10
                                is_overlay = False

                            # --- 2. TẠO LỚP ẢNH BỔ SUNG BẰNG FPDF ---
                            temp_dir = tempfile.gettempdir()
                            pdf_imgs = FPDF(orientation="L", unit="mm", format="A4")
                            
                            font_path = os.path.join(temp_dir, "Roboto-Bold.ttf")
                            if os.path.exists(font_path):
                                pdf_imgs.add_font("Roboto", style="B", fname=font_path)
                                
                            pdf_imgs.add_page()
                            # Bắt đầu vẽ từ tọa độ thông minh vừa tính toán
                            pdf_imgs.set_y(start_y) 
                            
                            try:
                                pdf_imgs.set_font("Roboto", style="B", size=10)
                            except:
                                pdf_imgs.set_font("Arial", style="B", size=10)
                                
                            pdf_imgs.cell(0, 6, "HÌNH ẢNH ĐỐI CHỨNG CHỨNG TỪ C/O GỐC (BỔ SUNG SAU):", new_x="LMARGIN", new_y="NEXT")
                            pdf_imgs.ln(4)
                            
                            img_pairs = [imgs_bo_sung[i:i + 2] for i in range(0, len(imgs_bo_sung), 2)]
                            for row_idx, pair in enumerate(img_pairs):
                                row_imgs_info = []
                                for idx_in_pair, img_file in enumerate(pair):
                                    eb_img = PILImage.open(img_file).convert('RGB')
                                    eb_path = os.path.join(temp_dir, f"temp_evidence_bs_{row_idx}_{idx_in_pair}.png")
                                    eb_img.save(eb_path, format="PNG")
                                    
                                    w_img_px, h_img_px = eb_img.size
                                    render_w = 125.0
                                    render_h = (h_img_px / w_img_px) * render_w
                                    if render_h > 75:
                                        render_h = 75.0
                                        render_w = (w_img_px / h_img_px) * render_h
                                        
                                    row_imgs_info.append({'path': eb_path, 'w': render_w, 'h': render_h})
                                    
                                max_row_h = max(item['h'] for item in row_imgs_info)
                                
                                # Nếu bị tràn trang thì tự động tạo trang mới
                                if pdf_imgs.get_y() + max_row_h + 6 > 195:
                                    pdf_imgs.add_page()
                                    
                                current_y = pdf_imgs.get_y() + 2
                                x_positions = [15.0, 148.0]
                                for col_idx, item in enumerate(row_imgs_info):
                                    x_pos = x_positions[col_idx]
                                    pdf_imgs.image(item['path'], x=x_pos, y=current_y, w=item['w'], h=item['h'])
                                    if os.path.exists(item['path']):
                                        os.remove(item['path'])
                                        
                                pdf_imgs.set_y(current_y + max_row_h + 5)
                            
                            imgs_pdf_buffer = io.BytesIO(pdf_imgs.output())
                            
                            # --- 3. GHÉP/XẾP CHỒNG (OVERLAY) LỚP ẢNH VÀO PDF GỐC ---
                            merger = pypdf.PdfWriter()
                            original_pdf = pypdf.PdfReader(io.BytesIO(file_bytes))
                            new_imgs_pdf = pypdf.PdfReader(imgs_pdf_buffer)
                            
                            if is_overlay:
                                # Chép nguyên xi các trang trước đó
                                for i in range(len(original_pdf.pages) - 1):
                                    merger.add_page(original_pdf.pages[i])
                                    
                                # Trộn lớp ảnh (trang 1) đè lên phần khoảng trắng của trang cuối file gốc
                                last_page = original_pdf.pages[-1]
                                last_page.merge_page(new_imgs_pdf.pages[0])
                                merger.add_page(last_page)
                                
                                # Nếu ảnh vượt qua 1 trang, chép tiếp các trang sau của file ảnh như bình thường
                                for i in range(1, len(new_imgs_pdf.pages)):
                                    merger.add_page(new_imgs_pdf.pages[i])
                            else:
                                # Nếu không đủ không gian, chép nối đuôi bình thường
                                for page in original_pdf.pages:
                                    merger.add_page(page)
                                for page in new_imgs_pdf.pages:
                                    merger.add_page(page)
                                    
                            final_buffer = io.BytesIO()
                            merger.write(final_buffer)
                            
                            st.session_state["pdf_ready_opt3"] = final_buffer.getvalue()
                            # [BỔ SUNG]: Lưu lại tên file gốc vào session_state để dùng sau khi rerun
                            st.session_state["ten_pdf_goc_opt3"] = file_pdf_da_ky.name 
                            st.rerun()

                        except ImportError:
                            st.error("⚠️ Hệ thống thiếu thư viện. Hãy đảm bảo đã cài đặt: `pip install pypdf pdfplumber`")
                        except Exception as e:
                            st.error(f"Lỗi khi ghép file: {e}")
            # ================= ĐOẠN CẦN BỔ SUNG LẠI =================
            # HIỂN THỊ NÚT TẢI SAU KHI XỬ LÝ (DÀNH CHO OPTION 3)
            # =========================================================
            # [CẬP NHẬT 2] - HIỂN THỊ NÚT TẢI VÀ NÚT TIẾP TỤC SONG SONG
            if st.session_state.get("pdf_ready_opt3"):
                #import datetime
                import os # Bổ sung thư viện xử lý đường dẫn
                
                st.success("✅ Đã xử lý phân tích và ghép ảnh thành công! Nút tải đã sẵn sàng.")
                
                # [BỔ SUNG]: Lấy tên file gốc và tạo tên mới
                ten_goc = st.session_state.get("ten_pdf_goc_opt3", "Ban_Giao_CO.pdf")
                ten_khong_duoi = os.path.splitext(ten_goc)[0]
                filename = f"{ten_khong_duoi}_Kem_Hinh.pdf"
                
                col_btn1, col_btn2 = st.columns([1, 1])
                
                with col_btn1:
                    st.download_button(
                        label="📥 BẤM VÀO ĐÂY ĐỂ TẢI FILE",
                        data=st.session_state["pdf_ready_opt3"],
                        file_name=filename,
                        mime="application/pdf",
                        type="primary",
                        use_container_width=True # Kéo dãn nút cho đẹp
                    )
                    
                with col_btn2:
                    if st.button("🔄 Tiếp Tục: Chèn ảnh vào file khác", use_container_width=True):
                        st.session_state["pdf_ready_opt3"] = None       # Ẩn nút tải
                        st.session_state["reset_uploader_key"] += 1     # Ép Streamlit làm mới file_uploader
                        st.rerun()                                      # Tải lại giao diện tức thì
            # =========================================================

        # ==========================================
        # HIỂN THỊ LƯỚI DATA VÀ KHUNG KÝ (CHỈ DÀNH CHO OPTION 1 & 2)
        # ==========================================
        if st.session_state.get("ocr_data") and che_do_tab3 in ["1️⃣ Quét PDF Mới & Chỉnh Sửa", "2️⃣ Tải File Excel Cũ Lên (Làm tiếp)"]:
            
            # --- [CẬP NHẬT MỚI]: BỘ CHỌN CÔNG TY ĐẠI LÝ ---
            st.markdown("<h4 style='margin-bottom: 5px; margin-top: 0px;'>🏢 Thông tin Đại lý nhận bàn giao</h4>", unsafe_allow_html=True)
            col_dl1, col_dl2 = st.columns([2, 3])
            
            cong_ty_dai_ly_chon = col_dl1.selectbox(
                "Chọn Công ty Đại lý:",
                options=["EXPEDITORS", "DSV AIR & SEA CO. LTD", "KUEHNE+NAGEL", "DHL", "KHÁC (Tự nhập tay)"]
            )
            
            if cong_ty_dai_ly_chon == "KHÁC (Tự nhập tay)":
                ten_cong_ty_dai_ly = col_dl2.text_input("Nhập tên công ty đại lý:", placeholder="VD: CÔNG TY LOGISTICS ABC...")
            else:
                ten_cong_ty_dai_ly = cong_ty_dai_ly_chon
                
            st.markdown("<hr style='margin: 10px 0px; border: 1px solid #ddd;'>", unsafe_allow_html=True)
            # ----------------------------------------------
            
            df_current = pd.DataFrame(st.session_state["ocr_data"])
            st.markdown(f"<h5 style='margin-bottom: 0px; margin-top: 0px;'>📝 Danh sách {len(df_current)} mã C/O</h5>", unsafe_allow_html=True)
            
            # [BỔ SUNG]: Hướng dẫn người dùng thao tác xóa dòng
            st.info("💡 **Mẹo:** Nhấp đúp chuột vào ô để sửa chữ. Để **XÓA DÒNG**, hãy click vào ô vuông đầu dòng bên trái cùng và nhấn phím **Delete** trên bàn phím (hoặc biểu tượng 🗑️).")
            
            edited_df = st.data_editor(
                df_current, 
                use_container_width=True, 
                num_rows="dynamic",
                disabled=["STT", "TÊN NGƯỜI NHẬN", "NGÀY NHẬN"],
                key="co_data_editor"
            )
            
            # =========================================================
            # [SỬA ĐỔI QUAN TRỌNG]: Tự động đánh lại STT liên tục từ 1, 2, 3...
            # =========================================================
            if not edited_df.empty:
                edited_df = edited_df.reset_index(drop=True)
                if 'STT' in edited_df.columns:
                    edited_df['STT'] = edited_df.index + 1
            
            # Lưu lại danh sách đã làm sạch STT vào session_state
            st.session_state["ocr_data"] = edited_df.to_dict('records')
            
            # [TẢI BẢN NHÁP EXCEL]
            
            
            import io
            excel_nhap_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_nhap_buffer, engine='xlsxwriter') as writer:
                workbook = writer.book
                worksheet = writer.sheets.setdefault('So_Giao_Nhan', workbook.add_worksheet('So_Giao_Nhan'))
                worksheet.write('A1', "CÔNG TY FORTUNATE HONGKONG VIỆT NAM")
                worksheet.write('A2', f"GỞI CÔNG TY: {ten_cong_ty_dai_ly}") # Gắn biến động vào file Excel
                edited_df.to_excel(writer, sheet_name='So_Giao_Nhan', index=False, startrow=3)
                worksheet.set_column('A:A', 10); worksheet.set_column('B:B', 35); worksheet.set_column('C:C', 30); worksheet.set_column('D:D', 25)
                
          #  import datetime
            st.download_button(
                label="📥 OPTION 1: Tải File Nháp (Chưa Ký) Để Lưu Trữ",
                data=excel_nhap_buffer.getvalue(),
                file_name=f"Ban_Giao_CO_{datetime.date.today().strftime('%d_%m_%Y')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            st.markdown("<hr style='margin: 10px 0px; border: 1px solid #ddd;'>", unsafe_allow_html=True)
            
            # [KHUNG KÝ TÊN VÀ XUẤT PDF]
            st.markdown("<h4 style='margin-bottom: 5px; margin-top: 0px;'>🤝 OPTION 2: Ký Nhận & Xuất Phiếu Bản Giao</h4>", unsafe_allow_html=True)

            st.markdown("<p style='margin-bottom: 0px; margin-top: 5px;'>📸 <b>Chụp/Upload hình ảnh C/O gốc làm bằng chứng (Không bắt buộc):</b></p>", unsafe_allow_html=True)
            imgs_bang_chung = st.file_uploader("Tải ảnh chụp tập C/O gốc (Thấy rõ mộc đỏ/số seri)", type=["jpg", "jpeg", "png"], accept_multiple_files=True, key="upload_evidence_co", label_visibility="collapsed")

            c_ten, c_ngay = st.columns(2)
            ten_dai_ly = c_ten.text_input("Tên nhân viên nhận (Đại lý):", placeholder="Nhập họ tên nhân viên nhận...")
            ngay_nhan = c_ngay.date_input("Ngày nhận:", value=datetime.date.today())
            
            st.markdown("<p style='margin-bottom: 0px; margin-top: 10px; font-weight: bold;'>Đại lý vẽ chữ ký xác nhận vào khung bên dưới:</p>", unsafe_allow_html=True)
            
            canvas_result = st_canvas(
                fill_color="rgba(255, 165, 0, 0.3)", 
                stroke_width=2.5, 
                stroke_color="#FF0000", 
                background_color="#f0f2f6", 
                height=150, 
                width=350, 
                drawing_mode="freedraw",
                return_image_data=True, key="canvas_ky_ten_2"
            )
            
            st.markdown("<div style='margin-top: -10px;'></div>", unsafe_allow_html=True)
            
            if st.button("🤝 KÝ VÀ TẢI FILE VỀ", type="primary"):
                if not ten_cong_ty_dai_ly.strip():
                    st.error("Vui lòng nhập Tên công ty đại lý.")
                elif not ten_dai_ly.strip():
                    st.error("Vui lòng nhập Tên nhân viên nhận của đại lý.")
                elif canvas_result.image_data is None or canvas_result.image_data.sum() == 0:
                    st.error("Đại lý chưa ký xác nhận vào khung.")
                else:
                    with st.spinner("Đang chèn chữ ký và xử lý File..."):
                        try:
                            from audit_logger import ghi_log_he_thong
                            chi_tiet_log = f"Bàn giao {len(edited_df)} C/O gốc cho đại lý: {ten_cong_ty_dai_ly} (Nhân viên: {ten_dai_ly})"
                            ghi_log_he_thong(db.pool, "QUAN_LY_CO", None, current_user, "BAN_GIAO_CO", chi_tiet_log)
                        except Exception as log_err:
                            pass
                        
                        img_data = canvas_result.image_data
                        img_rgba = Image.fromarray(img_data.astype('uint8'), 'RGBA')
                        img_bg = Image.new("RGBA", img_rgba.size, "WHITE")
                        img_bg.paste(img_rgba, mask=img_rgba)
                        img_rgb = img_bg.convert('RGB')
                        
                        img_buffer = io.BytesIO()
                        img_rgb.save(img_buffer, format="PNG")
                        img_buffer.seek(0)
                        
                        try:
                            from fpdf import FPDF
                        except ImportError:
                            st.error("⚠️ Thiếu thư viện PDF. Vui lòng thêm 'fpdf2' vào file requirements.txt")
                            st.stop()
                        
                        import urllib.request
                        import os
                        import tempfile
                        import base64

                        opener = urllib.request.build_opener()
                        opener.addheaders = [('User-agent', 'Mozilla/5.0')]
                        urllib.request.install_opener(opener)
                        
                        temp_dir = tempfile.gettempdir()
                        font_path = os.path.join(temp_dir, "Roboto-Regular.ttf")
                        font_bold_path = os.path.join(temp_dir, "Roboto-Bold.ttf")
                        
                        try:
                            if not os.path.exists(font_path):
                                urllib.request.urlretrieve("https://raw.githubusercontent.com/PolymerElements/font-roboto-local/master/fonts/roboto/Roboto-Regular.ttf", font_path)
                            if not os.path.exists(font_bold_path):
                                urllib.request.urlretrieve("https://raw.githubusercontent.com/PolymerElements/font-roboto-local/master/fonts/roboto/Roboto-Bold.ttf", font_bold_path)
                        except Exception as e:
                            st.error(f"Lỗi tải Font chữ: ({e})")
                            st.stop()
                            
                        pdf = FPDF(orientation="L", unit="mm", format="A4")
                        pdf.add_page()
                        
                        pdf.add_font("Roboto", style="", fname=font_path)
                        pdf.add_font("Roboto", style="B", fname=font_bold_path)
                        
                        pdf.set_font("Roboto", style="B", size=18)
                        pdf.cell(0, 10, "CÔNG TY FORTUNATE HONGKONG VIỆT NAM", align="C", new_x="LMARGIN", new_y="NEXT")
                        pdf.cell(0, 10, f"GỞI CÔNG TY: {ten_cong_ty_dai_ly}", align="C", new_x="LMARGIN", new_y="NEXT") # Gắn biến động vào PDF
                        pdf.ln(5)
                        
                        pdf.set_font("Roboto", style="B", size=11)
                        pdf.set_fill_color(220, 220, 220)
                        
                        col_widths = [20, 100, 80, 50] 
                        headers = ['STT', 'SỐ C/O BẢN GỐC', 'TÊN NGƯỜI NHẬN', 'NGÀY NHẬN']
                        
                        for i, h in enumerate(headers):
                            pdf.cell(col_widths[i], 8, h, border=1, align="C", fill=True)
                        pdf.ln(8)
                        
                        pdf.set_font("Roboto", style="", size=11)
                        ngay_format = ngay_nhan.strftime('%d/%m/%Y')
                        
                        for r_idx in range(len(edited_df)):
                            so_dong_con_lai = len(edited_df) - r_idx
                            space_needed = (so_dong_con_lai * 8) + 55 
                            
                            if so_dong_con_lai <= 3 and (pdf.get_y() + space_needed) > 200:
                                pdf.add_page()
                                pdf.set_font("Roboto", style="B", size=11)
                                for i, h in enumerate(headers):
                                    pdf.cell(col_widths[i], 8, h, border=1, align="C", fill=True)
                                pdf.ln(8)
                                pdf.set_font("Roboto", style="", size=11)
                            
                            pdf.cell(col_widths[0], 8, str(edited_df.iloc[r_idx, 0]).replace('.0', ''), border=1, align="C")
                            pdf.cell(col_widths[1], 8, str(edited_df.iloc[r_idx, 1]), border=1, align="C")
                            pdf.cell(col_widths[2], 8, ten_dai_ly, border=1, align="C") 
                            pdf.cell(col_widths[3], 8, ngay_format, border=1, align="C")
                            pdf.ln(8)
                            
                        if imgs_bang_chung:
                            try:
                                from PIL import Image as PILImage
                                if pdf.get_y() > 155:
                                    pdf.add_page()
                                else:
                                    pdf.ln(4)
                                    
                                pdf.set_font("Roboto", style="B", size=10)
                                pdf.cell(0, 6, "HÌNH ẢNH ĐỐI CHỨNG CHỨNG TỪ C/O GỐC THỰC TẾ KHI BÀN GIAO:", new_x="LMARGIN", new_y="NEXT")
                                
                                img_pairs = [imgs_bang_chung[i:i + 2] for i in range(0, len(imgs_bang_chung), 2)]
                                for row_idx, pair in enumerate(img_pairs):
                                    row_imgs_info = []
                                    for idx_in_pair, img_file in enumerate(pair):
                                        global_idx = row_idx * 2 + idx_in_pair
                                        eb_img = PILImage.open(img_file).convert('RGB')
                                        eb_path = os.path.join(temp_dir, f"temp_evidence_co_{global_idx}.png")
                                        eb_img.save(eb_path, format="PNG")
                                        
                                        w_img_px, h_img_px = eb_img.size
                                        render_w = 125.0
                                        render_h = (h_img_px / w_img_px) * render_w
                                        if render_h > 75:
                                            render_h = 75.0
                                            render_w = (w_img_px / h_img_px) * render_h
                                            
                                        row_imgs_info.append({'path': eb_path, 'w': render_w, 'h': render_h})
                                    
                                    max_row_h = max(item['h'] for item in row_imgs_info)
                                    if pdf.get_y() + max_row_h + 6 > 195:
                                        pdf.add_page()
                                        
                                    start_y = pdf.get_y() + 2
                                    x_positions = [15.0, 148.0]
                                    for col_idx, item in enumerate(row_imgs_info):
                                        x_pos = x_positions[col_idx]
                                        pdf.image(item['path'], x=x_pos, y=start_y, w=item['w'], h=item['h'])
                                        if os.path.exists(item['path']):
                                            os.remove(item['path'])
                                            
                                    pdf.set_y(start_y + max_row_h + 5)
                            except Exception as e_img:
                                print(f"Lỗi chèn ảnh bằng chứng: {e_img}")

                        if pdf.get_y() + 55 > 195:
                            pdf.add_page()
                        else:
                            pdf.ln(4)

                        pdf.set_font("Roboto", style="", size=9)
                        pdf.multi_cell(0, 5, "XÁC NHẬN: Đại lý cam kết đã kiểm tra kỹ, đếm và nhận đủ các bản C/O GỐC (có mộc đỏ) ghi trong danh sách trên. Mọi sự thất lạc sau thời điểm ký nhận này hoàn toàn thuộc trách nhiệm của Đại lý.", align="L")
                        pdf.ln(3)

                        pdf.set_font("Roboto", style="B", size=12)
                        pdf.cell(135, 8, "ĐẠI DIỆN BÀN GIAO", align="C")
                        pdf.cell(135, 8, "ĐẠI LÝ KÍ NHẬN - ĐÃ NHẬN ĐỦ", align="C", new_x="LMARGIN", new_y="NEXT")
                        
                        img_temp_path = os.path.join(tempfile.gettempdir(), "temp_signature.png")
                        img_rgb.save(img_temp_path, format="PNG")
                        current_y = pdf.get_y() + 2
                        pdf.image(img_temp_path, x=185, y=current_y, w=55) 
                        
                        if os.path.exists(img_temp_path):
                            os.remove(img_temp_path)
                            
                        pdf.ln(20)
                        pdf.set_font("Roboto", style="B", size=11)
                        pdf.cell(135, 6, "", align="C")
                        pdf.cell(135, 6, ten_dai_ly, align="C", new_x="LMARGIN", new_y="NEXT")
                        
                        pdf_bytes = pdf.output()
                        
                        ten_goc = st.session_state.get("ten_file_goc", "")
                        if ten_goc:
                            ten_file_xuat_ra = f"{os.path.splitext(ten_goc)[0]}_Da_Ky.pdf"
                        else:
                            ten_file_xuat_ra = f"Phieu_Ban_Giao_CO_{ten_dai_ly}_{datetime.date.today().strftime('%d_%m_%Y')}.pdf"
                        
                        st.success("✅ Đã tạo file thành công! Trình duyệt đang tự động tải xuống...")
                        
                        b64_pdf = base64.b64encode(pdf_bytes).decode()
                        href = f'''
                        <a id="auto_download_pdf" href="data:application/pdf;base64,{b64_pdf}" download="{ten_file_xuat_ra}" style="display:none;">Tải xuống</a>
                        <script>
                            setTimeout(function() {{
                                document.getElementById("auto_download_pdf").click();
                            }}, 500);
                        </script>
                        '''
                        st.components.v1.html(href, height=0)
                        st.markdown(f'<a href="data:application/pdf;base64,{b64_pdf}" download="{ten_file_xuat_ra}">🔗 Nhấp vào đây nếu trình duyệt không tự động tải xuống</a>', unsafe_allow_html=True)
            
            st.markdown("<div style='margin-top: 10px;'></div>", unsafe_allow_html=True)
            if st.button("🔄 Tải Lại Phiên Mới (Xóa Trắng)"):
                st.session_state["ocr_data"] = []
                st.session_state["ten_file_goc"] = ""
                keys_to_reset = ["upload_pdfs_ocr", "upload_excel_to_sign", "co_data_editor", "canvas_ky_ten_2"]
                for key in keys_to_reset:
                    if key in st.session_state:
                        del st.session_state[key]
                st.session_state["reset_uploader_key"] += 1
                st.rerun()

    vung_trich_xuat_va_ky_nhan()



    