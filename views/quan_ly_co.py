import streamlit as st
import pandas as pd
import datetime, time, re
import tempfile, os, io
import urllib.request

from utils_core import parse_money_input, tao_tieu_de_kem_nut_refresh
from co_manager import save_co_transaction, delete_co_transaction, get_don_gia_co_theo_khach_hang
from dotenv import load_dotenv
# Vẫn load file .env cho môi trường Local
load_dotenv(override=False) 

# --- HỆ THỐNG CACHE BỘ NHỚ ĐỆM ---
@st.cache_data(ttl=1800, show_spinner=False)
def get_cached_master_data(_db_instance, query, params=None):
    return _db_instance.execute_query(query, params)

def clear_master_cache():
    get_cached_master_data.clear()
# ---------------------------------

# 1. KIỂM TRA ĐĂNG NHẬP (Bắt buộc)
if 'username' not in st.session_state or not st.session_state['username']:
    st.error("⚠️ Phiên đăng nhập đã hết hạn hoặc bạn chưa đăng nhập. Vui lòng đăng nhập lại!")
    st.stop() # Dừng toàn bộ code phía dưới nếu chưa đăng nhập

# 2. LẤY CHÍNH XÁC USERNAME ĐÃ LOGIN
current_user = st.session_state['username']

# 3. KIỂM TRA KẾT NỐI DATABASE
db = st.session_state.get('db')
if not db:
    st.error("⚠️ Lỗi kết nối Cơ sở dữ liệu.")
    st.stop()

def get_idx(lst, val, default=0):
    try: return lst.index(val)
    except: return default

st.markdown("<h3 style='text-align: center; color: #0b5394;'>📄 PHÂN HỆ QUẢN LÝ CHỨNG TỪ C/O (XUẤT KHẨU)</h3>", unsafe_allow_html=True)
st.toast("✅ Đang load số liệu bạn vui lòng đợi nha:")
st.divider()

# Thêm tab OCR vào danh sách tabs
tab_khai_co, tab_file_co, tab_ocr_co, tab_quan_ly_co = st.tabs([
    "📋 KHAI BÁO C/O MỚI", 
    "📂 TẠO C/O TỪ FILE BỘ CÔNG THƯƠNG", 
    "🤖 TRÍCH XUẤT C/O TẠO FILE GIAO ĐẠI LÝ", 
    "🔍 DANH SÁCH & QUẢN LÝ"
])

# ==========================================
# TAB 1: KHAI BÁO C/O MỚI (TÍCH HỢP TỰ ĐỘNG ĐIỀN GIÁ)
# ==========================================
with tab_khai_co:
    st.markdown("#### 📥 Nhập Liệu Chứng Từ C/O Mới")
    @st.fragment
    def vung_thao_tac_declare_co():
        col_a, col_b, col_c = st.columns([5, 3, 2])
        
        sql_kh = "SELECT id, ten_khach_hang, ma_khach_hang FROM khach_hang ORDER BY ten_khach_hang ASC"
        df_kh = get_cached_master_data(db, sql_kh)
        dict_kh = {r['id']: f"[{r['ma_khach_hang']}] {r['ten_khach_hang']}" for _, r in df_kh.iterrows()} if not df_kh.empty else {}
        
        khach_hang_id = col_a.selectbox(
            "1. Chọn Khách Hàng*", options=list(dict_kh.keys()),
            format_func=lambda x: dict_kh[x],
            index=0)
        
        dict_tk = {}
        if khach_hang_id:
            sql_tk_xuat = "SELECT id, so_to_khai, ngay_khai FROM to_khai_hai_quan WHERE loai_to_khai = 'Xuat_Khau' AND khach_hang_id = %s ORDER BY id DESC"
            df_tk_xuat = get_cached_master_data(db, sql_tk_xuat, (khach_hang_id,))
            if isinstance(df_tk_xuat, pd.DataFrame) and not df_tk_xuat.empty:
                dict_tk = {r['id']: f"Số TK: {r['so_to_khai']} ({r['ngay_khai']})" for _, r in df_tk_xuat.iterrows()}
                
        to_khai_id = col_b.selectbox("2. Chọn Tờ Khai Của Khách*", options=list(dict_tk.keys()), format_func=lambda x: dict_tk[x]) if dict_tk else None
        
        if not to_khai_id and khach_hang_id:
            st.warning("⚠️ Khách hàng này chưa có tờ khai xuất khẩu nào trong hệ thống.")

        phan_loai_co = col_c.selectbox("3. Phân Loại Làm C/O", ["Thường", "Gấp", "Ghép"])
        
        gia_co_tu_dong = 0.0
        if khach_hang_id:
            gia_co_tu_dong = get_don_gia_co_theo_khach_hang(db.pool, khach_hang_id, phan_loai_co)
            
        st.divider()

        with st.form("form_khai_co", clear_on_submit=True):
            c1, c2, c3 = st.columns(3)
            form_co = c1.selectbox("Form C/O", ["", "Form EURO.1", "Form EURO.1UK", "Form AI", "Form EAV","Form AJ","Form AHK","Form RCEP","Form AK", "Form D","Form E","Form VK","Form VJ","Form S","Form VI"])
            so_co = c2.text_input("Số C/O*")
            # Lưu ý: GIỮ NGUYÊN kiểu dữ liệu datetime.date trả về, KHÔNG DÙNG strftime ở đây
            ngay_co = c3.date_input("Ngày Cấp C/O", value=datetime.date.today(), format="DD/MM/YYYY")
            
            c4, c5, c6 = st.columns(3)
            phi_co_val = f"{gia_co_tu_dong:,.0f}" if gia_co_tu_dong > 0 else ""
            
            phi_co = c4.text_input(f"Lệ Phí C/O (VNĐ)*", value=phi_co_val, placeholder="0", help="Hệ thống tự động đề xuất giá theo cấu hình. Có thể sửa tay.")
            phi_dvhq = c5.text_input("Phí DVHQ C/O (VNĐ)", value="", placeholder="0")
            so_hoa_don_co = c6.text_input("Số Hóa Đơn Phí C/O")
            
            ghi_chu_co = st.text_input("Ghi chú bổ sung", value=f"Phân loại C/O: {phan_loai_co}")
            
            if st.form_submit_button("💾 LƯU CHỨNG TỪ C/O", type="primary"):
                if not to_khai_id:
                    st.error("Vui lòng chọn tờ khai xuất khẩu hợp lệ.")
                elif not so_co:
                    st.error("Vui lòng nhập Số C/O.")
                else:
                    co_data = {
                        'to_khai_id': to_khai_id,
                        'form_co': form_co,
                        'so_co': so_co,
                        'ngay_co': ngay_co.strftime('%Y-%m-%d'),
                        'phi_co': parse_money_input(phi_co or 0),
                        'phi_dvhq': parse_money_input(phi_dvhq or 0),
                        'so_hoa_don_co': so_hoa_don_co,
                        'ghi_chu': ghi_chu_co
                    }
                    ok, msg = save_co_transaction(db.pool, co_data, None, current_user)
                    if ok:
                        clear_master_cache()
                        st.success("✅ Đã lưu chứng từ C/O thành công!")
                        st.session_state["select_co_action"] = None
                        time.sleep(1)
                        st.rerun()
                    else:
                        st.error(f"Lỗi: {msg}")
    vung_thao_tac_declare_co()

# ==========================================
# TAB 2: TẠO C/O TỪ FILE EXCEL
# ==========================================
with tab_file_co:
    st.markdown("#### 📥 Nhập Liệu Chứng Từ C/O Hàng Loạt Từ File Excel")
    st.info("Hỗ trợ đọc file danh sách C/O kết xuất từ phần mềm khai báo (VD: CO Fortuna09.xlsx). Hệ thống sẽ tự động đối chiếu Số tờ khai HQ để liên kết dữ liệu.")
    
    uploaded_file = st.file_uploader("📂 Chọn file Excel danh sách C/O", type=["xls", "xlsx"], key="upload_co_excel")
    
    if uploaded_file is not None:
        try:
            df_upload = pd.read_excel(uploaded_file)
            df_upload.dropna(how='all', inplace=True)
            df_upload.fillna('', inplace=True)
            
            required_cols = ['Số C/O', 'Số tờ khai HQ']
            missing = [c for c in required_cols if c not in df_upload.columns]
            
            if missing:
                st.error(f"❌ File Excel không đúng định dạng. Thiếu các cột: {', '.join(missing)}")
            else:
                sql_tk = "SELECT id, so_to_khai FROM to_khai_hai_quan"
                df_tk_sys = db.execute_query(sql_tk)
                
                dict_tk_sys = {}
                if isinstance(df_tk_sys, pd.DataFrame) and not df_tk_sys.empty:
                    for _, row in df_tk_sys.iterrows():
                        tk_str = str(row['so_to_khai']).strip()
                        if tk_str.endswith('.0'): tk_str = tk_str[:-2] 
                        dict_tk_sys[tk_str] = row['id']
                
                valid_records = []
                invalid_records = []
                
                for idx, row in df_upload.iterrows():
                    so_co = str(row.get('Số C/O', '')).strip()
                    if not so_co or so_co.lower() == 'nan':
                        continue 
                        
                    so_tk = str(row.get('Số tờ khai HQ', '')).strip()
                    if so_tk.endswith('.0'): so_tk = so_tk[:-2]
                    
                    tk_id = dict_tk_sys.get(so_tk)
                    
                    if not tk_id:
                        invalid_records.append({
                            "Số C/O": so_co,
                            "Số tờ khai HQ": so_tk,
                            "Lý do lỗi": "Không tìm thấy Số Tờ Khai này trong hệ thống."
                        })
                        continue
                        
                    phi_co = 0.0
                    raw_phi = row.get('PHI', '')
                    if str(raw_phi).strip():
                        try: phi_co = float(raw_phi)
                        except: pass
                        
                    ngay_co_val = datetime.date.today()
                    raw_ngay = row.get('Ngày cấp phép', '')
                    if str(raw_ngay).strip():
                        try:
                            ngay_co_val = pd.to_datetime(raw_ngay, dayfirst=True).date()
                        except: pass
                            
                    form_co = ""
                    if "VN-CN" in so_co.upper(): form_co = "Form E"
                    elif "VN-KR" in so_co.upper(): form_co = "Form VK"
                    elif "VN-JP" in so_co.upper(): form_co = "Form VJ"
                    elif "VN-CU" in so_co.upper(): form_co = "Form VN-CU"
                    
                    doanh_nghiep = str(row.get('Doanh nghiệp ký', '')).strip()
                    
                    valid_records.append({
                        "to_khai_id": tk_id,
                        "Số tờ khai": so_tk,
                        "Doanh nghiệp": doanh_nghiep,
                        "form_co": form_co,
                        "so_co": so_co,
                        "ngay_co": ngay_co_val,
                        "phi_co": phi_co,
                        "phi_dvhq": 0,
                        "so_hoa_don_co": "",
                        "ghi_chu": "Import hàng loạt từ Excel"
                    })
                
                if invalid_records:
                    st.warning(f"⚠️ Phát hiện {len(invalid_records)} dòng bị từ chối do chưa có Tờ Khai HQ trong phần mềm:")
                    st.dataframe(pd.DataFrame(invalid_records), use_container_width=True)
                    
                if valid_records:
                    st.success(f"✅ Đã chuẩn bị sẵn sàng {len(valid_records)} chứng từ C/O hợp lệ.")
                    df_valid = pd.DataFrame(valid_records)
                    
                    df_view = df_valid[['so_co', 'Số tờ khai', 'Doanh nghiệp', 'form_co', 'ngay_co', 'phi_co']].copy()
                    df_view['phi_co'] = df_view['phi_co'].apply(lambda x: f"{int(x):,}")
                    
                    st.dataframe(df_view, use_container_width=True)
                    
                    if st.button("🚀 XÁC NHẬN LƯU HÀNG LOẠT VÀO HỆ THỐNG", type="primary"):
                        with st.spinner("Đang kết nối Database và tiến hành lưu dữ liệu..."):
                            success_count = 0
                            error_msgs = []
                            
                            for rec in valid_records:
                                data_save = {
                                    'to_khai_id': rec['to_khai_id'],
                                    'form_co': rec['form_co'],
                                    'so_co': rec['so_co'],
                                    'ngay_co': rec['ngay_co'].strftime('%Y-%m-%d'),
                                    'phi_co': rec['phi_co'],
                                    'phi_dvhq': rec['phi_dvhq'],
                                    'so_hoa_don_co': rec['so_hoa_don_co'],
                                    'ghi_chu': rec['ghi_chu']
                                }
                                ok, msg = save_co_transaction(db.pool, data_save, None, current_user)
                                if ok: 
                                    success_count += 1
                                else: 
                                    error_msgs.append(f"- Lỗi ở Số C/O {rec['so_co']}: {msg}")
                            
                            clear_master_cache()
                            
                            if success_count > 0:
                                st.success(f"🎉 Hoàn tất! Đã lưu thành công {success_count}/{len(valid_records)} chứng từ C/O.")
                            if error_msgs:
                                st.error("❌ Các lỗi phát sinh trong quá trình lưu:")
                                for err in error_msgs: 
                                    st.write(err)
                                    
                            if success_count > 0:
                                time.sleep(1.5)
                                st.rerun()
                                    
        except Exception as e:
            st.error(f"❌ Xảy ra lỗi khi phân tích file Excel: {e}")

# ==========================================
# TAB 3: TRÍCH XUẤT C/O BẰNG AI & ĐẠI LÝ KÝ NHẬN BẢN GỐC
# ==========================================
with tab_ocr_co:
    st.markdown("#### 🤖 Trích Xuất Dữ Liệu C/O & Ký Nhận Bàn Giao")
    # 🚀 1. SỬ DỤNG @st.fragment ĐỂ CÔ LẬP RERUN (Tuyệt đối không giật trang khi gõ/ký tên)
    # 🚀 1. SỬ DỤNG @st.fragment ĐỂ CÔ LẬP RERUN (Tuyệt đối không giật trang khi gõ/ký tên)
    @st.fragment
    def vung_trich_xuat_va_ky_nhan():
        # --- BƯỚC 1: TIÊM CSS ÉP KHOẢNG CÁCH NGAY TỪ ĐẦU ---
        st.markdown("""
            <style>
                div[data-testid="stVerticalBlock"] > div { padding-bottom: 0.1rem !important; }
                div[data-testid="stHorizontalBlock"] > div { padding-bottom: 0.1rem !important; }
                .stTextInput > div, .stDateInput > div { margin-bottom: -15px; }
                h4, h5 { margin-top: 5px !important; margin-bottom: 5px !important; }
                div[role="radiogroup"] { margin-bottom: -10px !important; margin-top: -10px !important; }
                div[data-testid="stWebComponentContainer"] { touch-action: none !important; overscroll-behavior: none !important; }
                iframe[title="streamlit_drawable_canvas.st_canvas"] { touch-action: none !important; }
            </style>
        """, unsafe_allow_html=True)

        try:
            from streamlit_drawable_canvas import st_canvas
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.drawing.image import Image as OpenpyxlImage
            from PIL import Image
            import pdfplumber
            import easyocr
            import numpy as np
        except ImportError:
            st.error("⚠️ Server thiếu thư viện. Chạy lệnh: `pip install streamlit-drawable-canvas openpyxl pdfplumber easyocr pypdf fpdf2`")
            st.stop()

        # [CẬP NHẬT 1]: THÊM OPTION 3 LÊN TRÊN CÙNG
        che_do_tab3 = st.radio(
            "Lựa chọn luồng công việc:", 
            ["1️⃣ Quét PDF Mới & Chỉnh Sửa", "2️⃣ Tải File Excel Cũ Lên (Làm tiếp)", "3️⃣ Ghép Ảnh C/O Vào File Đã Ký"],
            horizontal=True
        )
        
        st.markdown("<hr style='margin: 10px 0px 5px 0px; border: 1px solid #e0e0e0;'>", unsafe_allow_html=True)
        
        if "ocr_data" not in st.session_state:
            st.session_state["ocr_data"] = []
        if "reset_uploader_key" not in st.session_state:
            st.session_state["reset_uploader_key"] = 0
        
        # ---------------------------------------------------------
        # CHẾ ĐỘ 1: QUÉT PDF MỚI VÀ CHỈNH SỬA TRỰC TIẾP
        # ---------------------------------------------------------
        if che_do_tab3 == "1️⃣ Quét PDF Mới & Chỉnh Sửa":
            current_pdf_key = f"upload_pdfs_ocr_{st.session_state['reset_uploader_key']}"
            uploaded_pdfs = st.file_uploader("📂 Kéo thả file PDF C/O vào đây", type=["pdf"], accept_multiple_files=True, key=current_pdf_key)
            
            if uploaded_pdfs:
                if st.button("🚀 Bắt Đầu Quét & Trích Xuất", type="primary"):
                    with st.spinner("⏳ Đang xử lý siêu tốc..."):
                        extracted_data = []
                        pattern = r'([A-Z0-9.\-\s]{0,20})(\d{2})\s*[-/|17lI\.]\s*(\d{2})\s*[-/|17lI\.]\s*([0-9A-Z\s]{5,15})'
                        
                        for pdf_file in uploaded_pdfs:
                            try:
                                with pdfplumber.open(pdf_file) as pdf:
                                    for page in pdf.pages:
                                        co_number = "Không nhận diện được"
                                        text_fast = (page.extract_text() or "").upper()
                                        match = re.search(pattern, text_fast)
                                        
                                        if not match:
                                            bounding_box = (page.width * 0.35, 0, page.width, page.height * 0.35)
                                            micro_crop = page.crop(bounding_box)
                                            pil_img = micro_crop.to_image(resolution=300).original
                                            
                                            try:
                                                import google.generativeai as genai
                                                if "api_keys" in st.secrets and "gemini" in st.secrets["api_keys"]:
                                                    api_key = st.secrets["api_keys"]["gemini"]
                                                elif "GEMINI_API_KEY" in st.secrets:
                                                    api_key = st.secrets["GEMINI_API_KEY"]
                                            except Exception:
                                                pass
                                                
                                            if not api_key:
                                                api_key = os.getenv("GEMINI_API_KEY")
                                                
                                            if api_key:
                                                genai.configure(api_key=api_key)
                                                model = genai.GenerativeModel('gemini-1.5-flash')
                                                prompt = "Đọc và trích xuất nguyên văn tất cả các đoạn văn bản có trong hình ảnh này. Không cần giải thích thêm."
                                                response = model.generate_content([prompt, pil_img])
                                                text_ai = (response.text or "").upper()
                                                match = re.search(pattern, text_ai)
                                            else:
                                                st.error("Lỗi: Không tìm thấy GEMINI_API_KEY trong hệ thống!")
                                                return None    
                                            
                                        if match:
                                            prefix_raw, dd, mm = match.group(1), match.group(2), match.group(3)
                                            raw_suffix = match.group(4)
                                            num_clean = re.sub(r'\s+', '', raw_suffix).replace('O', '0')
                                            
                                            m_strict = re.match(r'^(\d{5,8})([A-Z]?)', num_clean)
                                            if m_strict:
                                                digits = m_strict.group(1)
                                                letter = m_strict.group(2)
                                                if len(digits) >= 7 and digits[6] == '8':
                                                    digits = digits[:6] 
                                                    letter = 'B'        
                                                elif len(digits) > 6:
                                                    digits = digits[:6]
                                                if letter != 'B':
                                                    letter = ''
                                                num_clean = digits + letter
                                                
                                            suffix_clean = f"{dd}/{mm}/{num_clean}"
                                            prefix_clean = re.sub(r'[\s.\-]', '', prefix_raw)
                                            for noise in ['EUR1', 'EUR', 'NO', 'NUM']:
                                                prefix_clean = prefix_clean.replace(noise, '')
                                                
                                            country = "??"
                                            for c in ['DE', 'NL', 'SE', 'UK', 'KZ', 'K2', 'K7']:
                                                if c in prefix_clean:
                                                    country = c.replace('K2', 'KZ').replace('K7', 'KZ')
                                                    break
                                            if country == "??" and len(prefix_clean) >= 2: country = prefix_clean[-2:]
                                            if country == "??" and prefix_clean.endswith('E'): country = 'DE'
                                            co_number = f"VN-{country} {suffix_clean}" if country != "??" else f"VN-?? {suffix_clean}"
                                        
                                        extracted_data.append({"STT": len(extracted_data) + 1, "SỐ C/O BẢN GỐC": co_number, "KÝ NHẬN": "", "NGÀY NHẬN": ""})
                            except Exception as inner_e:
                                st.warning(f"Lỗi đọc {pdf_file.name}: {inner_e}")
                        
                        st.session_state["ocr_data"] = extracted_data
                        st.rerun()

        # ---------------------------------------------------------
        # CHẾ ĐỘ 2: TẢI LÊN FILE EXCEL (KHÔI PHỤC PHIÊN LÀM VIỆC)
        # ---------------------------------------------------------
        elif che_do_tab3 == "2️⃣ Tải File Excel Cũ Lên (Làm tiếp)":
            current_excel_key = f"upload_excel_to_sign_{st.session_state['reset_uploader_key']}"
            uploaded_excel = st.file_uploader("📂 Kéo thả file Excel (Bản chưa ký) vào đây", type=["xlsx", "xls"], key=current_excel_key)
            if uploaded_excel:
                try:
                    df_load = pd.read_excel(uploaded_excel, skiprows=3)
                    if 'SỐ C/O BẢN GỐC' in df_load.columns:
                        st.session_state["ocr_data"] = df_load.to_dict('records')
                        st.session_state["ten_file_goc"] = uploaded_excel.name
                        st.success("✅ Đã khôi phục dữ liệu từ file Excel thành công!")
                    else:
                        st.error("❌ File Excel không đúng định dạng chuẩn của phần mềm.")
                except Exception as e:
                    st.error(f"Lỗi đọc file Excel: {e}")

        # ---------------------------------------------------------
        # CHẾ ĐỘ 3: GHÉP ẢNH BỔ SUNG VÀO FILE PDF ĐÃ KÝ
        # ---------------------------------------------------------
        elif che_do_tab3 == "3️⃣ Ghép Ảnh C/O Vào File Đã Ký":
            st.info("💡 Mẹo: Ở công ty/cảng anh chỉ cần cho đại lý ký rồi tải file PDF chứa chữ ký về máy. Sau đó về nhà có mạng Wifi mạnh, anh tải ngược file PDF đó lên đây cùng với mớ hình ảnh C/O để hệ thống tự ghép ảnh vào đuôi PDF nhé!")
            
            col_up1, col_up2 = st.columns(2)
            
            # [CẬP NHẬT 1] - Đổi key tĩnh thành key động để dễ dàng xóa trắng
            current_pdf_key = f"upload_pdf_da_ky_{st.session_state['reset_uploader_key']}"
            current_img_key = f"upload_img_bo_sung_{st.session_state['reset_uploader_key']}"
            
            file_pdf_da_ky = col_up1.file_uploader("1️⃣ Tải lên Phiếu bàn giao ĐÃ KÝ (PDF)", type=["pdf"], key=current_pdf_key)
            imgs_bo_sung = col_up2.file_uploader("2️⃣ Tải lên hình ảnh C/O gốc bổ sung", type=["jpg", "jpeg", "png"], accept_multiple_files=True, key=current_img_key)
            
            if st.button("⚙️ BẮT ĐẦU GHÉP ẢNH VÀO PDF", type="primary"):
                if not file_pdf_da_ky:
                    st.error("⚠️ Vui lòng tải lên file PDF Phiếu bàn giao đã ký.")
                elif not imgs_bo_sung:
                    st.error("⚠️ Vui lòng tải lên ít nhất 1 hình ảnh C/O.")
                else:
                    with st.spinner("⏳ Đang tính toán không gian và nối ảnh, vui lòng chờ trong giây lát..."):
                        try:
                            import pypdf
                            from fpdf import FPDF
                            from PIL import Image as PILImage
                            import tempfile, os, io
                            import pdfplumber
                            
                            # --- 1. PHÂN TÍCH KHÔNG GIAN TRỐNG CỦA TRANG CUỐI BẰNG PDFPLUMBER ---
                            file_bytes = file_pdf_da_ky.getvalue()
                            lowest_y_mm = 0
                            try:
                                with pdfplumber.open(io.BytesIO(file_bytes)) as p_pdf:
                                    last_page = p_pdf.pages[-1]
                                    # Lấy tất cả đối tượng (chữ, hình ảnh chữ ký, đường kẻ) trên trang cuối
                                    all_objs = last_page.chars + last_page.images + last_page.rects + last_page.curves
                                    if all_objs:
                                        # Tìm tọa độ điểm thấp nhất (tính từ trên xuống dưới theo point)
                                        max_bottom_pt = max([obj['bottom'] for obj in all_objs if 'bottom' in obj], default=0)
                                        lowest_y_mm = max_bottom_pt / 2.83465 # Đổi từ point (72dpi) sang mm
                            except Exception as e:
                                print("Lỗi quét không gian PDF:", e)

                            # Chiều cao A4 ngang là 210mm. Một hàng ảnh cao tối đa 75mm + Tiêu đề = ~85mm
                            # Nếu điểm kết thúc nhỏ hơn 125mm -> Đủ không gian để chèn nối tiếp cùng trang
                            if 0 < lowest_y_mm < 125: 
                                start_y = lowest_y_mm + 15
                                is_overlay = True
                            else:
                                start_y = 10
                                is_overlay = False

                            # --- 2. TẠO LỚP ẢNH BỔ SUNG BẰNG FPDF ---
                            temp_dir = tempfile.gettempdir()
                            pdf_imgs = FPDF(orientation="L", unit="mm", format="A4")
                            
                            font_path = os.path.join(temp_dir, "Roboto-Bold.ttf")
                            if os.path.exists(font_path):
                                pdf_imgs.add_font("Roboto", style="B", fname=font_path)
                                
                            pdf_imgs.add_page()
                            # Bắt đầu vẽ từ tọa độ thông minh vừa tính toán
                            pdf_imgs.set_y(start_y) 
                            
                            try:
                                pdf_imgs.set_font("Roboto", style="B", size=10)
                            except:
                                pdf_imgs.set_font("Arial", style="B", size=10)
                                
                            pdf_imgs.cell(0, 6, "HÌNH ẢNH ĐỐI CHỨNG CHỨNG TỪ C/O GỐC (BỔ SUNG SAU):", new_x="LMARGIN", new_y="NEXT")
                            pdf_imgs.ln(4)
                            
                            img_pairs = [imgs_bo_sung[i:i + 2] for i in range(0, len(imgs_bo_sung), 2)]
                            for row_idx, pair in enumerate(img_pairs):
                                row_imgs_info = []
                                for idx_in_pair, img_file in enumerate(pair):
                                    eb_img = PILImage.open(img_file).convert('RGB')
                                    eb_path = os.path.join(temp_dir, f"temp_evidence_bs_{row_idx}_{idx_in_pair}.png")
                                    eb_img.save(eb_path, format="PNG")
                                    
                                    w_img_px, h_img_px = eb_img.size
                                    render_w = 125.0
                                    render_h = (h_img_px / w_img_px) * render_w
                                    if render_h > 75:
                                        render_h = 75.0
                                        render_w = (w_img_px / h_img_px) * render_h
                                        
                                    row_imgs_info.append({'path': eb_path, 'w': render_w, 'h': render_h})
                                    
                                max_row_h = max(item['h'] for item in row_imgs_info)
                                
                                # Nếu bị tràn trang thì tự động tạo trang mới
                                if pdf_imgs.get_y() + max_row_h + 6 > 195:
                                    pdf_imgs.add_page()
                                    
                                current_y = pdf_imgs.get_y() + 2
                                x_positions = [15.0, 148.0]
                                for col_idx, item in enumerate(row_imgs_info):
                                    x_pos = x_positions[col_idx]
                                    pdf_imgs.image(item['path'], x=x_pos, y=current_y, w=item['w'], h=item['h'])
                                    if os.path.exists(item['path']):
                                        os.remove(item['path'])
                                        
                                pdf_imgs.set_y(current_y + max_row_h + 5)
                            
                            imgs_pdf_buffer = io.BytesIO(pdf_imgs.output())
                            
                            # --- 3. GHÉP/XẾP CHỒNG (OVERLAY) LỚP ẢNH VÀO PDF GỐC ---
                            merger = pypdf.PdfWriter()
                            original_pdf = pypdf.PdfReader(io.BytesIO(file_bytes))
                            new_imgs_pdf = pypdf.PdfReader(imgs_pdf_buffer)
                            
                            if is_overlay:
                                # Chép nguyên xi các trang trước đó
                                for i in range(len(original_pdf.pages) - 1):
                                    merger.add_page(original_pdf.pages[i])
                                    
                                # Trộn lớp ảnh (trang 1) đè lên phần khoảng trắng của trang cuối file gốc
                                last_page = original_pdf.pages[-1]
                                last_page.merge_page(new_imgs_pdf.pages[0])
                                merger.add_page(last_page)
                                
                                # Nếu ảnh vượt qua 1 trang, chép tiếp các trang sau của file ảnh như bình thường
                                for i in range(1, len(new_imgs_pdf.pages)):
                                    merger.add_page(new_imgs_pdf.pages[i])
                            else:
                                # Nếu không đủ không gian, chép nối đuôi bình thường
                                for page in original_pdf.pages:
                                    merger.add_page(page)
                                for page in new_imgs_pdf.pages:
                                    merger.add_page(page)
                                    
                            final_buffer = io.BytesIO()
                            merger.write(final_buffer)
                            
                            st.session_state["pdf_ready_opt3"] = final_buffer.getvalue()
                            st.rerun()

                        except ImportError:
                            st.error("⚠️ Hệ thống thiếu thư viện. Hãy đảm bảo đã cài đặt: `pip install pypdf pdfplumber`")
                        except Exception as e:
                            st.error(f"Lỗi khi ghép file: {e}")
            # ================= ĐOẠN CẦN BỔ SUNG LẠI =================
            # HIỂN THỊ NÚT TẢI SAU KHI XỬ LÝ (DÀNH CHO OPTION 3)
            # =========================================================
            # [CẬP NHẬT 2] - HIỂN THỊ NÚT TẢI VÀ NÚT TIẾP TỤC SONG SONG
            if st.session_state.get("pdf_ready_opt3"):
                import datetime
                
                st.success("✅ Đã xử lý phân tích và ghép ảnh thành công! Nút tải đã sẵn sàng.")
                filename = f"Ban_Giao_CO_Kem_Hinh_{datetime.date.today().strftime('%d_%m_%Y')}.pdf"
                
                col_btn1, col_btn2 = st.columns([1, 1])
                
                with col_btn1:
                    st.download_button(
                        label="📥 BẤM VÀO ĐÂY ĐỂ TẢI FILE",
                        data=st.session_state["pdf_ready_opt3"],
                        file_name=filename,
                        mime="application/pdf",
                        type="primary",
                        use_container_width=True # Kéo dãn nút cho đẹp
                    )
                    
                with col_btn2:
                    if st.button("🔄 Tiếp Tục: Chèn ảnh vào file khác", use_container_width=True):
                        st.session_state["pdf_ready_opt3"] = None       # Ẩn nút tải
                        st.session_state["reset_uploader_key"] += 1     # Ép Streamlit làm mới file_uploader
                        st.rerun()                                      # Tải lại giao diện tức thì
            # =========================================================

        # ==========================================
        # HIỂN THỊ LƯỚI DATA VÀ KHUNG KÝ (CHỈ DÀNH CHO OPTION 1 & 2)
        # ==========================================
        if st.session_state.get("ocr_data") and che_do_tab3 in ["1️⃣ Quét PDF Mới & Chỉnh Sửa", "2️⃣ Tải File Excel Cũ Lên (Làm tiếp)"]:
            
            # --- [CẬP NHẬT MỚI]: BỘ CHỌN CÔNG TY ĐẠI LÝ ---
            st.markdown("<h4 style='margin-bottom: 5px; margin-top: 0px;'>🏢 Thông tin Đại lý nhận bàn giao</h4>", unsafe_allow_html=True)
            col_dl1, col_dl2 = st.columns([2, 3])
            
            cong_ty_dai_ly_chon = col_dl1.selectbox(
                "Chọn Công ty Đại lý:",
                options=["EXPEDITORS", "DSV AIR & SEA CO. LTD", "KUEHNE+NAGEL", "DHL", "KHÁC (Tự nhập tay)"]
            )
            
            if cong_ty_dai_ly_chon == "KHÁC (Tự nhập tay)":
                ten_cong_ty_dai_ly = col_dl2.text_input("Nhập tên công ty đại lý:", placeholder="VD: CÔNG TY LOGISTICS ABC...")
            else:
                ten_cong_ty_dai_ly = cong_ty_dai_ly_chon
                
            st.markdown("<hr style='margin: 10px 0px; border: 1px solid #ddd;'>", unsafe_allow_html=True)
            # ----------------------------------------------
            
            df_current = pd.DataFrame(st.session_state["ocr_data"])
            st.markdown(f"<h5 style='margin-bottom: 0px; margin-top: 0px;'>📝 Danh sách {len(df_current)} mã C/O (Nhấp đúp chuột để sửa)</h5>", unsafe_allow_html=True)
            
            edited_df = st.data_editor(
                df_current, 
                use_container_width=True, 
                num_rows="dynamic",
                disabled=["STT", "TÊN NGƯỜI NHẬN", "NGÀY NHẬN"],
                key="co_data_editor"
            )
            
            st.session_state["ocr_data"] = edited_df.to_dict('records')
            
            # [TẢI BẢN NHÁP EXCEL]
            import io
            excel_nhap_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_nhap_buffer, engine='xlsxwriter') as writer:
                workbook = writer.book
                worksheet = writer.sheets.setdefault('So_Giao_Nhan', workbook.add_worksheet('So_Giao_Nhan'))
                worksheet.write('A1', "CÔNG TY FORTUNATE HONGKONG VIỆT NAM")
                worksheet.write('A2', f"GỞI CÔNG TY: {ten_cong_ty_dai_ly}") # Gắn biến động vào file Excel
                edited_df.to_excel(writer, sheet_name='So_Giao_Nhan', index=False, startrow=3)
                worksheet.set_column('A:A', 10); worksheet.set_column('B:B', 35); worksheet.set_column('C:C', 30); worksheet.set_column('D:D', 25)
                
            import datetime
            st.download_button(
                label="📥 OPTION 1: Tải File Nháp (Chưa Ký) Để Lưu Trữ",
                data=excel_nhap_buffer.getvalue(),
                file_name=f"Ban_Giao_CO_{datetime.date.today().strftime('%d_%m_%Y')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

            st.markdown("<hr style='margin: 10px 0px; border: 1px solid #ddd;'>", unsafe_allow_html=True)
            
            # [KHUNG KÝ TÊN VÀ XUẤT PDF]
            st.markdown("<h4 style='margin-bottom: 5px; margin-top: 0px;'>🤝 OPTION 2: Ký Nhận & Xuất Phiếu Bản Giao</h4>", unsafe_allow_html=True)

            st.markdown("<p style='margin-bottom: 0px; margin-top: 5px;'>📸 <b>Chụp/Upload hình ảnh C/O gốc làm bằng chứng (Không bắt buộc):</b></p>", unsafe_allow_html=True)
            imgs_bang_chung = st.file_uploader("Tải ảnh chụp tập C/O gốc (Thấy rõ mộc đỏ/số seri)", type=["jpg", "jpeg", "png"], accept_multiple_files=True, key="upload_evidence_co", label_visibility="collapsed")

            c_ten, c_ngay = st.columns(2)
            ten_dai_ly = c_ten.text_input("Tên nhân viên nhận (Đại lý):", placeholder="Nhập họ tên nhân viên nhận...")
            ngay_nhan = c_ngay.date_input("Ngày nhận:", value=datetime.date.today())
            
            st.markdown("<p style='margin-bottom: 0px; margin-top: 10px; font-weight: bold;'>Đại lý vẽ chữ ký xác nhận vào khung bên dưới:</p>", unsafe_allow_html=True)
            
            canvas_result = st_canvas(
                fill_color="rgba(255, 165, 0, 0.3)", 
                stroke_width=2.5, 
                stroke_color="#FF0000", 
                background_color="#f0f2f6", 
                height=150, 
                width=350, 
                drawing_mode="freedraw",
                return_image_data=True, key="canvas_ky_ten_2"
            )
            
            st.markdown("<div style='margin-top: -10px;'></div>", unsafe_allow_html=True)
            
            if st.button("🤝 KÝ VÀ TẢI FILE VỀ", type="primary"):
                if not ten_cong_ty_dai_ly.strip():
                    st.error("Vui lòng nhập Tên công ty đại lý.")
                elif not ten_dai_ly.strip():
                    st.error("Vui lòng nhập Tên nhân viên nhận của đại lý.")
                elif canvas_result.image_data is None or canvas_result.image_data.sum() == 0:
                    st.error("Đại lý chưa ký xác nhận vào khung.")
                else:
                    with st.spinner("Đang chèn chữ ký và xử lý File..."):
                        try:
                            from audit_logger import ghi_log_he_thong
                            chi_tiet_log = f"Bàn giao {len(edited_df)} C/O gốc cho đại lý: {ten_cong_ty_dai_ly} (Nhân viên: {ten_dai_ly})"
                            ghi_log_he_thong(db.pool, "QUAN_LY_CO", None, current_user, "BAN_GIAO_CO", chi_tiet_log)
                        except Exception as log_err:
                            pass
                        
                        img_data = canvas_result.image_data
                        img_rgba = Image.fromarray(img_data.astype('uint8'), 'RGBA')
                        img_bg = Image.new("RGBA", img_rgba.size, "WHITE")
                        img_bg.paste(img_rgba, mask=img_rgba)
                        img_rgb = img_bg.convert('RGB')
                        
                        img_buffer = io.BytesIO()
                        img_rgb.save(img_buffer, format="PNG")
                        img_buffer.seek(0)
                        
                        try:
                            from fpdf import FPDF
                        except ImportError:
                            st.error("⚠️ Thiếu thư viện PDF. Vui lòng thêm 'fpdf2' vào file requirements.txt")
                            st.stop()
                        
                        import urllib.request
                        import os
                        import tempfile
                        import base64

                        opener = urllib.request.build_opener()
                        opener.addheaders = [('User-agent', 'Mozilla/5.0')]
                        urllib.request.install_opener(opener)
                        
                        temp_dir = tempfile.gettempdir()
                        font_path = os.path.join(temp_dir, "Roboto-Regular.ttf")
                        font_bold_path = os.path.join(temp_dir, "Roboto-Bold.ttf")
                        
                        try:
                            if not os.path.exists(font_path):
                                urllib.request.urlretrieve("https://raw.githubusercontent.com/PolymerElements/font-roboto-local/master/fonts/roboto/Roboto-Regular.ttf", font_path)
                            if not os.path.exists(font_bold_path):
                                urllib.request.urlretrieve("https://raw.githubusercontent.com/PolymerElements/font-roboto-local/master/fonts/roboto/Roboto-Bold.ttf", font_bold_path)
                        except Exception as e:
                            st.error(f"Lỗi tải Font chữ: ({e})")
                            st.stop()
                            
                        pdf = FPDF(orientation="L", unit="mm", format="A4")
                        pdf.add_page()
                        
                        pdf.add_font("Roboto", style="", fname=font_path)
                        pdf.add_font("Roboto", style="B", fname=font_bold_path)
                        
                        pdf.set_font("Roboto", style="B", size=18)
                        pdf.cell(0, 10, "CÔNG TY FORTUNATE HONGKONG VIỆT NAM", align="C", new_x="LMARGIN", new_y="NEXT")
                        pdf.cell(0, 10, f"GỞI CÔNG TY: {ten_cong_ty_dai_ly}", align="C", new_x="LMARGIN", new_y="NEXT") # Gắn biến động vào PDF
                        pdf.ln(5)
                        
                        pdf.set_font("Roboto", style="B", size=11)
                        pdf.set_fill_color(220, 220, 220)
                        
                        col_widths = [20, 100, 80, 50] 
                        headers = ['STT', 'SỐ C/O BẢN GỐC', 'TÊN NGƯỜI NHẬN', 'NGÀY NHẬN']
                        
                        for i, h in enumerate(headers):
                            pdf.cell(col_widths[i], 8, h, border=1, align="C", fill=True)
                        pdf.ln(8)
                        
                        pdf.set_font("Roboto", style="", size=11)
                        ngay_format = ngay_nhan.strftime('%d/%m/%Y')
                        
                        for r_idx in range(len(edited_df)):
                            so_dong_con_lai = len(edited_df) - r_idx
                            space_needed = (so_dong_con_lai * 8) + 55 
                            
                            if so_dong_con_lai <= 3 and (pdf.get_y() + space_needed) > 200:
                                pdf.add_page()
                                pdf.set_font("Roboto", style="B", size=11)
                                for i, h in enumerate(headers):
                                    pdf.cell(col_widths[i], 8, h, border=1, align="C", fill=True)
                                pdf.ln(8)
                                pdf.set_font("Roboto", style="", size=11)
                            
                            pdf.cell(col_widths[0], 8, str(edited_df.iloc[r_idx, 0]).replace('.0', ''), border=1, align="C")
                            pdf.cell(col_widths[1], 8, str(edited_df.iloc[r_idx, 1]), border=1, align="C")
                            pdf.cell(col_widths[2], 8, ten_dai_ly, border=1, align="C") 
                            pdf.cell(col_widths[3], 8, ngay_format, border=1, align="C")
                            pdf.ln(8)
                            
                        if imgs_bang_chung:
                            try:
                                from PIL import Image as PILImage
                                if pdf.get_y() > 155:
                                    pdf.add_page()
                                else:
                                    pdf.ln(4)
                                    
                                pdf.set_font("Roboto", style="B", size=10)
                                pdf.cell(0, 6, "HÌNH ẢNH ĐỐI CHỨNG CHỨNG TỪ C/O GỐC THỰC TẾ KHI BÀN GIAO:", new_x="LMARGIN", new_y="NEXT")
                                
                                img_pairs = [imgs_bang_chung[i:i + 2] for i in range(0, len(imgs_bang_chung), 2)]
                                for row_idx, pair in enumerate(img_pairs):
                                    row_imgs_info = []
                                    for idx_in_pair, img_file in enumerate(pair):
                                        global_idx = row_idx * 2 + idx_in_pair
                                        eb_img = PILImage.open(img_file).convert('RGB')
                                        eb_path = os.path.join(temp_dir, f"temp_evidence_co_{global_idx}.png")
                                        eb_img.save(eb_path, format="PNG")
                                        
                                        w_img_px, h_img_px = eb_img.size
                                        render_w = 125.0
                                        render_h = (h_img_px / w_img_px) * render_w
                                        if render_h > 75:
                                            render_h = 75.0
                                            render_w = (w_img_px / h_img_px) * render_h
                                            
                                        row_imgs_info.append({'path': eb_path, 'w': render_w, 'h': render_h})
                                    
                                    max_row_h = max(item['h'] for item in row_imgs_info)
                                    if pdf.get_y() + max_row_h + 6 > 195:
                                        pdf.add_page()
                                        
                                    start_y = pdf.get_y() + 2
                                    x_positions = [15.0, 148.0]
                                    for col_idx, item in enumerate(row_imgs_info):
                                        x_pos = x_positions[col_idx]
                                        pdf.image(item['path'], x=x_pos, y=start_y, w=item['w'], h=item['h'])
                                        if os.path.exists(item['path']):
                                            os.remove(item['path'])
                                            
                                    pdf.set_y(start_y + max_row_h + 5)
                            except Exception as e_img:
                                print(f"Lỗi chèn ảnh bằng chứng: {e_img}")

                        if pdf.get_y() + 55 > 195:
                            pdf.add_page()
                        else:
                            pdf.ln(4)

                        pdf.set_font("Roboto", style="", size=9)
                        pdf.multi_cell(0, 5, "XÁC NHẬN: Đại lý cam kết đã kiểm tra kỹ, đếm và nhận đủ các bản C/O GỐC (có mộc đỏ) ghi trong danh sách trên. Mọi sự thất lạc sau thời điểm ký nhận này hoàn toàn thuộc trách nhiệm của Đại lý.", align="L")
                        pdf.ln(3)

                        pdf.set_font("Roboto", style="B", size=12)
                        pdf.cell(135, 8, "ĐẠI DIỆN BÀN GIAO", align="C")
                        pdf.cell(135, 8, "ĐẠI LÝ KÍ NHẬN - ĐÃ NHẬN ĐỦ", align="C", new_x="LMARGIN", new_y="NEXT")
                        
                        img_temp_path = os.path.join(tempfile.gettempdir(), "temp_signature.png")
                        img_rgb.save(img_temp_path, format="PNG")
                        current_y = pdf.get_y() + 2
                        pdf.image(img_temp_path, x=185, y=current_y, w=55) 
                        
                        if os.path.exists(img_temp_path):
                            os.remove(img_temp_path)
                            
                        pdf.ln(20)
                        pdf.set_font("Roboto", style="B", size=11)
                        pdf.cell(135, 6, "", align="C")
                        pdf.cell(135, 6, ten_dai_ly, align="C", new_x="LMARGIN", new_y="NEXT")
                        
                        pdf_bytes = pdf.output()
                        
                        ten_goc = st.session_state.get("ten_file_goc", "")
                        if ten_goc:
                            ten_file_xuat_ra = f"{os.path.splitext(ten_goc)[0]}_Da_Ky.pdf"
                        else:
                            ten_file_xuat_ra = f"Phieu_Ban_Giao_CO_{ten_dai_ly}_{datetime.date.today().strftime('%d_%m_%Y')}.pdf"
                        
                        st.success("✅ Đã tạo file thành công! Trình duyệt đang tự động tải xuống...")
                        
                        b64_pdf = base64.b64encode(pdf_bytes).decode()
                        href = f'''
                        <a id="auto_download_pdf" href="data:application/pdf;base64,{b64_pdf}" download="{ten_file_xuat_ra}" style="display:none;">Tải xuống</a>
                        <script>
                            setTimeout(function() {{
                                document.getElementById("auto_download_pdf").click();
                            }}, 500);
                        </script>
                        '''
                        st.components.v1.html(href, height=0)
                        st.markdown(f'<a href="data:application/pdf;base64,{b64_pdf}" download="{ten_file_xuat_ra}">🔗 Nhấp vào đây nếu trình duyệt không tự động tải xuống</a>', unsafe_allow_html=True)
            
            st.markdown("<div style='margin-top: 10px;'></div>", unsafe_allow_html=True)
            if st.button("🔄 Tải Lại Phiên Mới (Xóa Trắng)"):
                st.session_state["ocr_data"] = []
                st.session_state["ten_file_goc"] = ""
                keys_to_reset = ["upload_pdfs_ocr", "upload_excel_to_sign", "co_data_editor", "canvas_ky_ten_2"]
                for key in keys_to_reset:
                    if key in st.session_state:
                        del st.session_state[key]
                st.session_state["reset_uploader_key"] += 1
                st.rerun()

    vung_trich_xuat_va_ky_nhan()


# ==========================================
# TAB 4: DANH SÁCH & QUẢN LÝ (SỬA / XÓA)
# ==========================================
with tab_quan_ly_co:
    tao_tieu_de_kem_nut_refresh("🔍 Danh sách Chứng từ C/O", "ref_tab_ds_co")
    @st.fragment
    def vung_thao_tac_edit_delete_co():
        col_f1, col_f2 = st.columns(2)
        today = datetime.date.today()
        co_tu_ngay = col_f1.date_input("Từ ngày", value=today.replace(day=1),format= "DD/MM/YYYY", key="co_tu_ngay")
        co_den_ngay = col_f2.date_input("Đến ngày", value=today,format= "DD/MM/YYYY", key="co_den_ngay")
        
        sql_ds_co = """
            SELECT co.id, co.to_khai_id, co.form_co, co.so_co, co.ngay_co, 
                co.phi_co, co.phi_dvhq, co.so_hoa_don_co, co.ghi_chu, 
                tk.so_to_khai, kh.ten_khach_hang, kh.id AS khach_hang_id
            FROM to_khai_co co
            JOIN to_khai_hai_quan tk ON co.to_khai_id = tk.id
            JOIN khach_hang kh ON tk.khach_hang_id = kh.id
            WHERE co.ngay_co BETWEEN %s AND %s
            ORDER BY co.ngay_co DESC, co.id DESC
        """
        df_co = db.execute_query(sql_ds_co, (co_tu_ngay.strftime('%Y-%m-%d'), co_den_ngay.strftime('%Y-%m-%d')))
        
        if isinstance(df_co, pd.DataFrame) and not df_co.empty:
            df_co['phi_co'] = pd.to_numeric(df_co['phi_co'], errors='coerce').fillna(0)
            df_co['phi_dvhq'] = pd.to_numeric(df_co['phi_dvhq'], errors='coerce').fillna(0)
            df_co['so_hoa_don_co'] = df_co['so_hoa_don_co'].fillna('').apply(lambda x: str(x).strip() if str(x).strip().lower() != 'nan' else '')
            df_co['ghi_chu'] = df_co['ghi_chu'].fillna('').apply(lambda x: str(x).strip() if str(x).strip().lower() != 'nan' else '')

            df_co_view = df_co[['id', 'form_co', 'so_co', 'ngay_co', 'so_to_khai', 'ten_khach_hang', 'phi_co', 'phi_dvhq', 'so_hoa_don_co', 'ghi_chu']].copy()
            # [BỔ SUNG]: Format cột ngay_co hiển thị theo dạng dd/mm/yyyy
            df_co_view['ngay_co'] = pd.to_datetime(df_co_view['ngay_co'], errors='coerce').dt.strftime('%d/%m/%Y').fillna('')
            df_co_view['phi_co'] = df_co_view['phi_co'].apply(lambda x: f"{int(x):,}" if pd.notnull(x) else "0")
            df_co_view['phi_dvhq'] = df_co_view['phi_dvhq'].apply(lambda x: f"{int(x):,}" if pd.notnull(x) else "0")
            st.dataframe(df_co_view, use_container_width=True, hide_index=True)
            
            st.divider()
            st.markdown("#### 🛠️ Thao Tác Quản Lý C/O (Sửa / Xóa)")
            
            dict_co_edit = {row['id']: f"Form: {row['form_co']} | Số C/O: {row['so_co']} (TK: {row['so_to_khai']} - {row['ten_khach_hang']})" for _, row in df_co.iterrows()}
            
            selected_co_id = st.selectbox(
                "📌 Chọn chứng từ C/O để sửa hoặc xóa:",
                options=list(dict_co_edit.keys()),
                format_func=lambda x: dict_co_edit[x],
                index=None,
                placeholder="-- Vui lòng click chọn 1 chứng từ C/O --",
                key="select_co_action"
            )
            
            if selected_co_id is not None:
                co_info = df_co[df_co['id'] == selected_co_id].iloc[0]
                
                def get_safe_val(key, default=""):
                    val = co_info.get(key)
                    if pd.isna(val) or str(val).strip() == "" or str(val).strip().lower() == 'nan':
                        return default
                    return str(val).strip()

                def get_safe_float(key, default=0.0):
                    val = co_info.get(key)
                    if pd.isna(val) or str(val).strip() == "" or str(val).strip().lower() == 'nan':
                        return default
                    try: return float(val)
                    except: return default

                st.markdown(f"Đang thao tác với Số C/O: **{get_safe_val('so_co')}**")
                action_mode_co = st.radio("Hành động:", ["✏️ Sửa C/O", "🗑️ Xóa C/O"], horizontal=True, key="radio_co_action")
                
                if action_mode_co == "🗑️ Xóa C/O":
                    st.warning(f"⚠️ Bạn có chắc chắn muốn xóa vĩnh viễn chứng từ C/O **{get_safe_val('so_co')}**?")
                    if st.button("Xác Nhận Xóa C/O", type="primary"):
                        ok, msg = delete_co_transaction(db.pool, selected_co_id, current_user)
                        if ok:
                            clear_master_cache()
                            st.success("✅ Đã xóa chứng từ C/O thành công!")
                            if "select_co_action" in st.session_state:
                                del st.session_state["select_co_action"]
                            st.rerun()
                        else:
                            st.error(f"Lỗi: {msg}")
                else:
                    with st.form(f"form_edit_co_{selected_co_id}", clear_on_submit=False):
                        sql_tk_edit = "SELECT id, so_to_khai FROM to_khai_hai_quan WHERE khach_hang_id = %s"
                        df_tk_edit = get_cached_master_data(db, sql_tk_edit, (co_info['khach_hang_id'],))
                        # [FIX LỖI]: Kiểm tra isinstance để đảm bảo nó là DataFrame trước khi gọi .empty
                        if isinstance(df_tk_edit, pd.DataFrame) and not df_tk_edit.empty:
                            dict_tk_edit = {r['id']: f"Số TK: {r['so_to_khai']}" for _, r in df_tk_edit.iterrows()}
                        else:
                            dict_tk_edit = {co_info['to_khai_id']: co_info['so_to_khai']}
                        
                        e_to_khai_id = st.selectbox("Tờ Khai Xuất Khẩu Liên Kết", options=list(dict_tk_edit.keys()), index=get_idx(list(dict_tk_edit.keys()), co_info['to_khai_id']), format_func=lambda x: dict_tk_edit[x])
                            
                        ec1, ec2, ec3 = st.columns(3)
                        e_form_co = ec1.text_input("Loại Form C/O", value=get_safe_val('form_co'))
                        e_so_co = ec2.text_input("Số C/O*", value=get_safe_val('so_co'))
                        
                        raw_ngay_co = co_info.get('ngay_co')
                        default_ngay_co = pd.to_datetime(raw_ngay_co).date() if pd.notna(raw_ngay_co) and str(raw_ngay_co).strip().lower() != 'nan' else datetime.date.today()
                        e_ngay_co = ec3.date_input("Ngày Cấp C/O", value=default_ngay_co)
                        
                        def fmt(val): 
                            if pd.isna(val) or val == "" or str(val).strip().lower() == 'nan': return ""
                            try:
                                num = float(val)
                                return f"{int(num):,}" if num > 0 else ""
                            except: return ""
                        
                        ec4, ec5, ec6 = st.columns(3)
                        e_phi_co = ec4.text_input("Lệ Phí C/O (VNĐ)*", value=fmt(get_safe_float('phi_co')), placeholder="0")
                        e_phi_dvhq = ec5.text_input("Phí DVHQ C/O (VNĐ)", value=fmt(get_safe_float('phi_dvhq')), placeholder="0")
                        e_so_hoa_don_co = ec6.text_input("Số Hóa Đơn Phí C/O", value=get_safe_val('so_hoa_don_co'))
                        
                        e_ghi_chu = st.text_input("Ghi chú bổ sung", value=get_safe_val('ghi_chu'))
                        
                        if st.form_submit_button("💾 LƯU THAY ĐỔI C/O", type="primary"):
                            if not e_so_co:
                                st.error("Vui lòng nhập số C/O.")
                            else:
                                co_update_data = {
                                    'to_khai_id': e_to_khai_id,
                                    'form_co': e_form_co,
                                    'so_co': e_so_co,
                                    'ngay_co': e_ngay_co.strftime('%Y-%m-%d'),
                                    'phi_co': parse_money_input(e_phi_co or 0),
                                    'phi_dvhq': parse_money_input(e_phi_dvhq or 0),
                                    'so_hoa_don_co': e_so_hoa_don_co,
                                    'ghi_chu': e_ghi_chu
                                }
                                ok, msg = save_co_transaction(db.pool, co_update_data, selected_co_id, current_user)
                                if ok:
                                    clear_master_cache()
                                    st.success("✅ Cập nhật chứng từ C/O thành công!")
                                    if "select_co_action" in st.session_state:
                                        del st.session_state["select_co_action"]
                                    st.rerun()
                                else:
                                    st.error(f"Lỗi: {msg}")
            else:
                st.info("👆 Vui lòng chọn một chứng từ C/O từ danh sách bên trên để tiến hành sửa hoặc xóa.")
        else:
            st.info("📭 Không có chứng từ C/O nào trong khoảng thời gian này.")
    vung_thao_tac_edit_delete_co()